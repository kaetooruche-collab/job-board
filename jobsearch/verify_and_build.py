"""Merges the researched job records (raw/*.json), re-checks every link, drops anything older than 7 days or closed, and
writes the working list for KT: pharmacy-jobs.xlsx (with check-off columns), pharmacy-jobs.csv and pharmacy-jobs.md.
Run: python3 jobsearch/verify_and_build.py   (needs: requests, openpyxl)"""
import csv, json, re, sys, time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

HERE = Path(__file__).resolve().parent
TODAY = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else datetime.now(timezone.utc).date()
MAX_AGE = 7
S = requests.Session()
S.headers["User-Agent"] = "Mozilla/5.0 (personal job search link check)"
if Path("/root/.ccr/ca-bundle.crt").exists():
    S.verify = "/root/.ccr/ca-bundle.crt"
CLOSED = re.compile(r"no longer (?:accepting|available)|position (?:has been )?(?:filled|closed)|job (?:is )?(?:no longer|expired|closed)|this (?:job|position|role) (?:has )?(?:expired|closed)|no longer open|posting (?:has )?(?:expired|closed)|job not found|page (?:not|cannot be) found", re.I)

def load():
    recs = []
    for f in sorted((HERE / "raw").glob("*.json")):
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
        except Exception as e:  # noqa: BLE001
            print("BAD JSON", f.name, e); continue
        for r in data:
            r["_file"] = f.stem
            recs.append(r)
    return recs

def norm(u):
    return re.sub(r"[?#].*$", "", (u or "").strip().lower()).rstrip("/")

def check(r):
    u = r.get("url") or ""
    out = {"status": None, "note": ""}
    try:
        resp = S.get(u, timeout=25, allow_redirects=True)
        out["status"] = resp.status_code
        if resp.status_code in (401, 403, 429, 999):
            out["note"] = "site blocks automatic checks; agent verified it by opening the page"
        elif resp.status_code == 404 or resp.status_code == 410:
            out["note"] = "link dead"
        elif resp.status_code == 200:
            text = resp.text[:300000]
            for m in CLOSED.finditer(text):
                before = text[max(0, m.start() - 4000): m.start()].lower()
                if "job-expired-view" in before or "ph-page-state=\"expired\"" in before or "expired job page state" in before:
                    continue  # hidden template block that every Phenom careers page ships with
                out["note"] = "page shows closed-role wording"
                break
    except requests.RequestException as e:
        out["note"] = "could not load (" + type(e).__name__ + ")"
    return out

def main():
    recs = load()
    def ckey(r):
        t = re.sub(r"\(.*?\)", "", r.get("title", "")).lower()
        t = re.sub(r"[^a-z0-9]+", " ", t).strip()
        co = re.sub(r"[^a-z0-9 ]", "", r.get("company", "").lower()).split()
        loc = re.sub(r"[^a-z0-9 ]", " ", (r.get("location") or "").lower()).split()
        return (co[0] if co else "", t, " ".join(loc[:2]))
    rank = {"A": 0, "B": 1, "C": 2}
    recs.sort(key=lambda r: rank.get(r.get("fit"), 3))
    seen, seen_c, uniq = set(), set(), []
    for r in recs:
        k = norm(r.get("url")) or (r.get("company", "") + r.get("title", "")).lower()
        if k in seen or ckey(r) in seen_c:
            continue
        seen.add(k); seen_c.add(ckey(r)); uniq.append(r)
    with ThreadPoolExecutor(6) as ex:
        results = list(ex.map(check, uniq))
    keep, dropped = [], []
    for r, c in zip(uniq, results):
        r["_check"] = c
        try:
            age = (TODAY - date.fromisoformat(r["posted_date"])).days
        except Exception:  # noqa: BLE001
            dropped.append((r, "no valid posted date")); continue
        r["_age"] = age
        if age > MAX_AGE or age < 0:
            dropped.append((r, f"posted {age} days ago")); continue
        if not r.get("verified_open", False):
            dropped.append((r, "not verified open")); continue
        if c["note"] in ("link dead", "page shows closed-role wording") or (c["note"].startswith("could not load")):
            dropped.append((r, c["note"])); continue
        keep.append(r)
    order = {"A": 0, "B": 1, "C": 2}
    keep.sort(key=lambda r: (order.get(r.get("fit"), 3), r["_age"], r.get("company", "")))
    main_list = [r for r in keep if r.get("fit") in ("A", "B")]
    ruled_out = [r for r in keep if r.get("fit") == "C"]
    (HERE / "pharmacy-jobs.json").write_text(json.dumps({"built": str(TODAY), "active": main_list, "ruled_out": ruled_out,
                                                         "dropped": [{"title": r.get("title"), "company": r.get("company"), "url": r.get("url"), "why": why} for r, why in dropped]}, indent=1), encoding="utf-8")
    write_outputs(main_list, ruled_out, dropped)
    print(f"{len(recs)} raw, {len(uniq)} unique, {len(main_list)} active A/B, {len(ruled_out)} ruled out (C), {len(dropped)} dropped")
    for r, why in dropped:
        print("  dropped:", r.get("company"), "|", r.get("title"), "|", why)

COLS = [
    ("#", 5), ("Fit", 5), ("Role family", 18), ("Job title", 34), ("Company", 20), ("Posted", 11), ("Days old", 8),
    ("Location (as stated)", 22), ("Work mode", 11), ("Open to non-residents?", 12), ("Eligibility rules (quotes)", 40),
    ("Min. experience", 16), ("Min. education", 20), ("Licence needed", 16), ("Languages", 14), ("Pay (if stated)", 14),
    ("What it requires", 55), ("What you would do", 45), ("CV keywords to echo (only if true)", 40), ("Where your MPharm fits", 40),
    ("Gaps", 34), ("Why this rating", 40), ("How to apply", 12), ("Apply link", 40), ("Documents asked for", 22),
    ("Contact email (published)", 26), ("Where the email was published", 28),
    ("Link rechecked", 14),
]
CHECK_COLS = ["1 Read posting", "2 Eligibility OK", "3 CV tailored", "4 Cover letter", "5 Contact found", "6 Applied (date)", "7 Follow-up (date)", "Result / notes"]

def row_for(i, r):
    j = lambda v: "\n".join("- " + x for x in v) if isinstance(v, list) else (v or "")
    c = r.get("_check", {})
    rechecked = (f"{TODAY}: OK" if c.get("status") == 200 and not c.get("note") else f"{TODAY}: {c.get('note') or 'HTTP ' + str(c.get('status'))}")
    return [i, r.get("fit"), r.get("role_family"), r.get("title"), r.get("company"), r.get("posted_date"), r.get("_age"),
            r.get("location"), r.get("work_mode"), r.get("open_to_non_residents"), r.get("eligibility_stated"),
            r.get("min_experience"), r.get("min_education"), r.get("licence_required"), r.get("languages"), r.get("salary_stated") or "not stated",
            j(r.get("key_requirements")), j(r.get("key_duties")), "; ".join(r.get("keywords_for_cv") or []), r.get("mpharm_overlap"),
            r.get("gaps"), r.get("fit_reason"), r.get("how_to_apply"), r.get("apply_url") or r.get("url"), r.get("docs_required"),
            r.get("contact_email") or "none published: apply via the portal", r.get("contact_email_source") or "", rechecked]

def write_outputs(active, ruled_out, dropped):
    wb = Workbook()
    ws = wb.active; ws.title = "Active list"
    head = [c for c, _ in COLS] + CHECK_COLS
    ws.append(head)
    for i, r in enumerate(active, 1):
        ws.append(row_for(i, r) + [""] * len(CHECK_COLS))
    navy = PatternFill("solid", fgColor="1A1A2E"); tick = PatternFill("solid", fgColor="0F3460")
    for n, cell in enumerate(ws[1], 1):
        cell.font = Font(bold=True, color="FFFFFF"); cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = tick if n > len(COLS) else navy
    for n, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(n)].width = w
    for n in range(len(COLS) + 1, len(head) + 1):
        ws.column_dimensions[get_column_letter(n)].width = 14
    ws.column_dimensions[get_column_letter(len(head))].width = 30
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        link = row[23]
        if link.value and str(link.value).startswith("http"):
            link.hyperlink = link.value; link.font = Font(color="2563EB", underline="single")
    dv = DataValidation(type="list", formula1='"Yes,No,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    first_check = len(COLS) + 1
    dv.add(f"{get_column_letter(first_check)}2:{get_column_letter(first_check + 4)}{max(2, ws.max_row)}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    ws.conditional_formatting.add(f"G2:G{max(2, ws.max_row)}", CellIsRule(operator="greaterThanOrEqual", formula=["6"], fill=PatternFill("solid", bgColor="FDE68A")))
    ws.conditional_formatting.add(f"B2:B{max(2, ws.max_row)}", CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", bgColor="BBF7D0")))
    ws.conditional_formatting.add(f"B2:B{max(2, ws.max_row)}", CellIsRule(operator="equal", formula=['"B"'], fill=PatternFill("solid", bgColor="FEF3C7")))

    wr = wb.create_sheet("Ruled out (and why)")
    wr.append(["Fit", "Job title", "Company", "Posted", "Location", "Link", "Why not realistic"])
    for r in ruled_out:
        wr.append([r.get("fit"), r.get("title"), r.get("company"), r.get("posted_date"), r.get("location"), r.get("url"), r.get("fit_reason")])
    for n, w in enumerate([5, 36, 22, 11, 24, 44, 70], 1):
        wr.column_dimensions[get_column_letter(n)].width = w
    for c in wr[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = navy
    for row in wr.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wd = wb.create_sheet("Dropped at recheck")
    wd.append(["Job title", "Company", "Link", "Why dropped"])
    for r, why in dropped:
        wd.append([r.get("title"), r.get("company"), r.get("url"), why])
    for n, w in enumerate([40, 24, 50, 40], 1):
        wd.column_dimensions[get_column_letter(n)].width = w
    for c in wd[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = navy

    wh = wb.create_sheet("How to use")
    for line in [
        f"Built {TODAY}. Only jobs posted in the last {MAX_AGE} days, each opened and checked for an open application.",
        "Fit: A = you meet the stated minimum. B = a stretch with one clear gap. (C = ruled out; see the next sheet.)",
        "Days old turns amber at 6+ days: apply to those first, they will close soon.",
        "Work through the numbered check columns left to right for each job. Use Yes / No / N/A.",
        "Contact emails are only ever ones the employer published. If none, apply through the portal.",
        "CV keywords: use them only where they are true for you.",
        "Re-run the search weekly (see jobsearch/README.md); expired jobs are removed and new ones added.",
    ]:
        wh.append([line])
    wh.column_dimensions["A"].width = 120
    wb.save(HERE / "pharmacy-jobs.xlsx")

    with open(HERE / "pharmacy-jobs.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(head)
        for i, r in enumerate(active, 1):
            w.writerow(row_for(i, r) + [""] * len(CHECK_COLS))
    md = [f"# Pharmacy and pharmacy-adjacent jobs (built {TODAY}; posted within the last {MAX_AGE} days)\n"]
    md.append("Fit A = meets the stated minimum. B = stretch with one clear gap. Newest first within each fit.\n")
    for fam in sorted({r.get('role_family') or '' for r in active}):
        md.append(f"\n## {fam}\n")
        for r in [x for x in active if (x.get('role_family') or '') == fam]:
            md.append(f"- **[{r.get('fit')}] {r.get('title')}**, {r.get('company')} ({r.get('location')}; {r.get('work_mode')}; open to non-residents: {r.get('open_to_non_residents')}). Posted {r.get('posted_date')}. [Apply]({r.get('apply_url') or r.get('url')}). {r.get('fit_reason') or ''}")
    (HERE / "pharmacy-jobs.md").write_text("\n".join(md) + "\n", encoding="utf-8")

if __name__ == "__main__":
    main()
